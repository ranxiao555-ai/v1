#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import warnings

warnings.filterwarnings("ignore", category=DeprecationWarning)

import cgi
import io
import json
import os
import random
import sqlite3
import sys
import traceback
from datetime import date, datetime, timedelta
from http import HTTPStatus
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
UPLOAD_DIR = BASE_DIR / "uploads"
EXPORT_DIR = BASE_DIR / "exports"
STATIC_DIR = BASE_DIR / "static"
DB_PATH = DATA_DIR / "sales_report.db"

REQUIRED_SALES_FIELDS = [
    "日期",
    "业务员",
    "客户名称",
    "商品编码",
    "商品名称",
    "销售数量",
    "单价",
    "销售金额",
]

BASE_IMPORT_CONFIG = {
    "salespeople": {
        "title": "业务员基础资料",
        "headers": ["业务员编码", "业务员姓名", "所属区域", "所属部门", "入职日期", "状态"],
    },
    "products": {
        "title": "商品基础资料",
        "headers": ["商品编码", "商品名称", "商品规格", "商品分类", "是否重点商品", "状态"],
    },
    "goldProducts": {
        "title": "重点商品维护",
        "headers": ["商品编码", "商品名称", "金砖商品目标", "每日目标"],
    },
    "salesTargets": {
        "title": "销售目标维护",
        "headers": ["业务员", "巅峰目标", "每日目标"],
    },
}

SALESPEOPLE = [
    ("张三", "华东区", 1),
    ("李四", "华南区", 1),
    ("王五", "华北区", 1),
    ("赵六", "西南区", 1),
    ("孙七", "华中区", 1),
]

PRODUCT_CATEGORIES = ["肉类", "海鲜", "速冻", "调味", "饮品", "粮油", "日配", "休闲"]
KEY_PRODUCT_NAMES = ["牛肉卷", "羊肉卷", "虾滑", "鸡胸肉", "鳕鱼排", "火锅丸子", "肥牛片", "培根", "鸡翅中", "酸汤肥牛料"]


def month_days(month_text):
    year, month = map(int, month_text.split("-"))
    if month == 12:
        return (date(year + 1, 1, 1) - date(year, month, 1)).days
    return (date(year, month + 1, 1) - date(year, month, 1)).days


def month_end(month_text):
    year, month = map(int, month_text.split("-"))
    if month == 12:
        return (date(year + 1, 1, 1) - timedelta(days=1)).isoformat()
    return (date(year, month + 1, 1) - timedelta(days=1)).isoformat()


def ensure_dirs():
    for path in [DATA_DIR, UPLOAD_DIR, EXPORT_DIR, STATIC_DIR]:
        path.mkdir(parents=True, exist_ok=True)


def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db(seed=False):
    ensure_dirs()
    export_import_template()
    export_base_templates()
    with db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS salespeople (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                code TEXT,
                name TEXT NOT NULL UNIQUE,
                region TEXT NOT NULL,
                department TEXT NOT NULL DEFAULT '',
                entry_date TEXT,
                active INTEGER NOT NULL DEFAULT 1,
                pending_confirm INTEGER NOT NULL DEFAULT 0
            );

            CREATE TABLE IF NOT EXISTS products (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                code TEXT NOT NULL UNIQUE,
                name TEXT NOT NULL,
                spec TEXT NOT NULL DEFAULT '',
                category TEXT NOT NULL,
                is_key INTEGER NOT NULL DEFAULT 0,
                gold_target REAL NOT NULL DEFAULT 0,
                gold_daily_target REAL NOT NULL DEFAULT 0,
                status TEXT NOT NULL DEFAULT '启用',
                pending_confirm INTEGER NOT NULL DEFAULT 0
            );

            CREATE TABLE IF NOT EXISTS salesperson_targets (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                month TEXT NOT NULL,
                salesperson_name TEXT NOT NULL,
                peak_target REAL NOT NULL DEFAULT 0,
                daily_target REAL NOT NULL DEFAULT 0,
                UNIQUE(month, salesperson_name)
            );

            CREATE TABLE IF NOT EXISTS gold_targets (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                month TEXT NOT NULL,
                salesperson_name TEXT NOT NULL,
                target_qty REAL NOT NULL DEFAULT 0,
                UNIQUE(month, salesperson_name)
            );

            CREATE TABLE IF NOT EXISTS targets (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                month TEXT NOT NULL,
                salesperson_name TEXT NOT NULL,
                product_code TEXT NOT NULL,
                product_name TEXT NOT NULL,
                target_qty REAL NOT NULL,
                UNIQUE(month, salesperson_name, product_code)
            );

            CREATE TABLE IF NOT EXISTS sales (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                sale_date TEXT NOT NULL,
                salesperson_name TEXT NOT NULL,
                customer_name TEXT NOT NULL,
                product_code TEXT NOT NULL,
                product_name TEXT NOT NULL,
                quantity REAL NOT NULL,
                unit_price REAL NOT NULL,
                amount REAL NOT NULL,
                uploaded_at TEXT NOT NULL,
                source_file TEXT
            );

            CREATE TABLE IF NOT EXISTS pending_items (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                item_type TEXT NOT NULL,
                ref_key TEXT NOT NULL,
                name TEXT NOT NULL,
                payload TEXT,
                status TEXT NOT NULL DEFAULT '待确认',
                created_at TEXT NOT NULL
            );
            """
        )
        ensure_schema(conn)
        if seed:
            seed_demo_data(conn)


def ensure_schema(conn):
    def columns(table):
        return {row["name"] for row in conn.execute(f"PRAGMA table_info({table})")}

    salespeople_cols = columns("salespeople")
    if "code" not in salespeople_cols:
        conn.execute("ALTER TABLE salespeople ADD COLUMN code TEXT")
    if "department" not in salespeople_cols:
        conn.execute("ALTER TABLE salespeople ADD COLUMN department TEXT NOT NULL DEFAULT ''")
    if "entry_date" not in salespeople_cols:
        conn.execute("ALTER TABLE salespeople ADD COLUMN entry_date TEXT")
    if "pending_confirm" not in salespeople_cols:
        conn.execute("ALTER TABLE salespeople ADD COLUMN pending_confirm INTEGER NOT NULL DEFAULT 0")

    product_cols = columns("products")
    if "spec" not in product_cols:
        conn.execute("ALTER TABLE products ADD COLUMN spec TEXT NOT NULL DEFAULT ''")
    if "gold_target" not in product_cols:
        conn.execute("ALTER TABLE products ADD COLUMN gold_target REAL NOT NULL DEFAULT 0")
    if "gold_daily_target" not in product_cols:
        conn.execute("ALTER TABLE products ADD COLUMN gold_daily_target REAL NOT NULL DEFAULT 0")
    if "status" not in product_cols:
        conn.execute("ALTER TABLE products ADD COLUMN status TEXT NOT NULL DEFAULT '启用'")
    if "pending_confirm" not in product_cols:
        conn.execute("ALTER TABLE products ADD COLUMN pending_confirm INTEGER NOT NULL DEFAULT 0")
    conn.commit()


def seed_demo_data(conn):
    conn.executescript(
        """
        DELETE FROM sales;
        DELETE FROM targets;
        DELETE FROM gold_targets;
        DELETE FROM salesperson_targets;
        DELETE FROM pending_items;
        DELETE FROM products;
        DELETE FROM salespeople;
        """
    )
    for row in SALESPEOPLE:
        code = f"YW{SALESPEOPLE.index(row) + 1:03d}"
        conn.execute(
            "INSERT INTO salespeople(code, name, region, department, entry_date, active, pending_confirm) VALUES (?, ?, ?, ?, ?, ?, 0)",
            (code, row[0], row[1], "销售部", date.today().replace(month=1, day=1).isoformat(), row[2]),
        )

    products = []
    for i in range(1, 101):
        code = f"SP{i:03d}"
        if i <= 10:
            name = KEY_PRODUCT_NAMES[i - 1]
            is_key = 1
        else:
            name = f"普通商品{i:03d}"
            is_key = 0
        category = PRODUCT_CATEGORIES[(i - 1) % len(PRODUCT_CATEGORIES)]
        products.append((code, name, category, is_key))
        gold_target = random.randint(300, 800) if is_key else 0
        gold_daily_target = round(gold_target / month_days(date.today().strftime("%Y-%m")), 2) if is_key else 0
        conn.execute(
            """
            INSERT INTO products(code, name, spec, category, is_key, gold_target, gold_daily_target, status, pending_confirm)
            VALUES (?, ?, ?, ?, ?, ?, ?, '启用', 0)
            """,
            (code, name, "标准规格", category, is_key, gold_target, gold_daily_target),
        )

    today = date.today()
    month = today.strftime("%Y-%m")
    random.seed(20260620)

    for salesperson, _, _ in SALESPEOPLE:
        conn.execute(
            """
            INSERT INTO salesperson_targets(month, salesperson_name, peak_target, daily_target)
            VALUES (?, ?, ?, ?)
            """,
            (month, salesperson, random.randint(180000, 280000), random.randint(6500, 10000)),
        )
        conn.execute(
            """
            INSERT INTO gold_targets(month, salesperson_name, target_qty)
            VALUES (?, ?, ?)
            """,
            (month, salesperson, random.randint(1200, 1900)),
        )
        for code, name, _, is_key in products[:10]:
            target = random.randint(260, 620)
            conn.execute(
                """
                INSERT INTO targets(month, salesperson_name, product_code, product_name, target_qty)
                VALUES (?, ?, ?, ?, ?)
                """,
                (month, salesperson, code, name, target),
            )

    customers = [f"客户{i:03d}" for i in range(1, 81)]
    price_map = {code: round(random.uniform(8, 98), 2) for code, _, _, _ in products}
    uploaded_at = datetime.now().isoformat(timespec="seconds")

    for offset in range(29, -1, -1):
        current = today - timedelta(days=offset)
        rows = random.randint(45, 85)
        for _ in range(rows):
            salesperson = random.choice(SALESPEOPLE)[0]
            product = random.choice(products)
            code, name, _, is_key = product
            qty = random.randint(2, 35) if is_key else random.randint(1, 18)
            unit_price = price_map[code]
            amount = round(qty * unit_price, 2)
            conn.execute(
                """
                INSERT INTO sales(
                    sale_date, salesperson_name, customer_name, product_code, product_name,
                    quantity, unit_price, amount, uploaded_at, source_file
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    current.isoformat(),
                    salesperson,
                    random.choice(customers),
                    code,
                    name,
                    qty,
                    unit_price,
                    amount,
                    uploaded_at,
                    "系统模拟数据",
                ),
            )
    conn.commit()
    export_demo_upload_file(conn)


def export_demo_upload_file(conn):
    rows = conn.execute(
        """
        SELECT sale_date AS 日期, salesperson_name AS 业务员, customer_name AS 客户名称,
               product_code AS 商品编码, product_name AS 商品名称, quantity AS 销售数量,
               unit_price AS 单价, amount AS 销售金额
        FROM sales
        ORDER BY sale_date, id
        """
    ).fetchall()
    write_xlsx(EXPORT_DIR / "模拟销售明细_最近30天.xlsx", REQUIRED_SALES_FIELDS, [dict(r) for r in rows])


def export_import_template():
    template_rows = [
        {
            "日期": date.today().isoformat(),
            "业务员": "张三",
            "客户名称": "示例客户",
            "商品编码": "SP001",
            "商品名称": "示例商品",
            "销售数量": 10,
            "单价": 25.5,
            "销售金额": "",
        }
    ]
    write_xlsx(EXPORT_DIR / "销售明细导入模板.xlsx", REQUIRED_SALES_FIELDS, template_rows)


def export_base_templates():
    samples = {
        "salespeople": [{"业务员编码": "YW001", "业务员姓名": "张三", "所属区域": "华东区", "所属部门": "销售部", "入职日期": date.today().isoformat(), "状态": "启用"}],
        "products": [{"商品编码": "SP001", "商品名称": "牛肉卷", "商品规格": "500g/袋", "商品分类": "肉类", "是否重点商品": "是", "状态": "启用"}],
        "goldProducts": [{"商品编码": "SP001", "商品名称": "牛肉卷", "金砖商品目标": 300, "每日目标": 10}],
        "salesTargets": [{"业务员": "张三", "巅峰目标": 200000, "每日目标": 8000}],
    }
    for key, config in BASE_IMPORT_CONFIG.items():
        write_xlsx(EXPORT_DIR / f"{config['title']}导入模板.xlsx", config["headers"], samples[key])


def read_excel_dicts(file_bytes, filename, required_headers):
    suffix = Path(filename or "").suffix.lower()
    if suffix not in (".xlsx", ".xls"):
        return [], ["仅支持 .xlsx/.xls 文件。"]
    try:
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception:
        if suffix == ".xls":
            return [], [".xls 文件格式较旧，当前本地环境无法解析；请在 Excel 中另存为 .xlsx 后导入。"]
        raise
    ws = wb.active
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
    missing = [header for header in required_headers if header not in headers]
    if missing:
        return [], [f"缺少必要字段：{', '.join(missing)}"]
    idx = {name: headers.index(name) for name in required_headers}
    rows = []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        data = {name: row[idx[name]] for name in required_headers}
        data["_row_no"] = row_no
        rows.append(data)
    return rows, []


def truthy(value):
    return str(value or "").strip() in ("是", "启用", "1", "true", "True", "Y", "y")


def status_active(value):
    return 0 if str(value or "").strip() in ("停用", "离职", "否", "0") else 1


def parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ["%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"]:
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def to_float(value):
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).replace(",", "").strip())


def read_sales_excel(file_bytes):
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
    missing = [field for field in REQUIRED_SALES_FIELDS if field not in headers]
    if missing:
        return None, [f"缺少必要字段：{', '.join(missing)}"]

    idx = {name: headers.index(name) for name in REQUIRED_SALES_FIELDS}
    rows = []
    errors = []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue
        try:
            sale_date = parse_date(row[idx["日期"]])
            salesperson = str(row[idx["业务员"]] or "").strip()
            customer = str(row[idx["客户名称"]] or "").strip()
            product_code = str(row[idx["商品编码"]] or "").strip()
            product_name = str(row[idx["商品名称"]] or "").strip()
            qty = to_float(row[idx["销售数量"]])
            unit_price = to_float(row[idx["单价"]])
            amount_cell = row[idx["销售金额"]]
            amount = to_float(amount_cell) if amount_cell not in (None, "") else round(qty * unit_price, 2)
            if not sale_date:
                raise ValueError("日期格式无效")
            if not salesperson or not customer or not product_code or not product_name:
                raise ValueError("业务员、客户名称、商品编码、商品名称不能为空")
            rows.append(
                {
                    "sale_date": sale_date,
                    "salesperson_name": salesperson,
                    "customer_name": customer,
                    "product_code": product_code,
                    "product_name": product_name,
                    "quantity": qty,
                    "unit_price": unit_price,
                    "amount": round(amount, 2),
                }
            )
        except Exception as exc:
            errors.append(f"第 {row_no} 行：{exc}")
    return rows, errors


def write_xlsx(path, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([row.get(header, "") for header in headers])
    for col in range(1, len(headers) + 1):
        max_len = max(len(str(ws.cell(r, col).value or "")) for r in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 12), 28)
    ws.freeze_panes = "A2"
    wb.save(path)


def query_rows(sql, params=None):
    with db() as conn:
        return [dict(r) for r in conn.execute(sql, params or {}).fetchall()]


def one(sql, params=None):
    with db() as conn:
        row = conn.execute(sql, params or {}).fetchone()
        return dict(row) if row else {}


def date_range(period, selected_date):
    base = datetime.strptime(selected_date, "%Y-%m-%d").date() if selected_date else date.today()
    if period == "week":
        start = base - timedelta(days=base.weekday())
        end = start + timedelta(days=6)
    elif period == "month":
        start = base.replace(day=1)
        if start.month == 12:
            end = date(start.year + 1, 1, 1) - timedelta(days=1)
        else:
            end = date(start.year, start.month + 1, 1) - timedelta(days=1)
    else:
        start = end = base
    return start.isoformat(), end.isoformat()


def filters_from(query):
    return {k: v[0] for k, v in parse_qs(query).items()}


def salesperson_stats(params):
    start, end = date_range(params.get("period", "day"), params.get("date"))
    extra = ""
    sql_params = {"start": start, "end": end}
    if params.get("salesperson"):
        extra += " AND salesperson_name = :salesperson"
        sql_params["salesperson"] = params["salesperson"]
    rows = query_rows(
        f"""
        SELECT salesperson_name AS 业务员,
               ROUND(SUM(amount), 2) AS 销售金额,
               SUM(quantity) AS 销售数量,
               COUNT(DISTINCT customer_name) AS 成交客户数,
               COUNT(DISTINCT product_code) AS 销售商品数
        FROM sales
        WHERE sale_date BETWEEN :start AND :end {extra}
        GROUP BY salesperson_name
        ORDER BY 销售金额 DESC
        """,
        sql_params,
    )
    for i, row in enumerate(rows, 1):
        row["排名"] = i
    return rows


def product_stats(params):
    start, end = date_range(params.get("period", "day"), params.get("date"))
    extra = ""
    sql_params = {"start": start, "end": end}
    if params.get("product"):
        extra += " AND (product_code = :product OR product_name LIKE :product_like)"
        sql_params["product"] = params["product"]
        sql_params["product_like"] = f"%{params['product']}%"
    if params.get("keyOnly") == "1":
        extra += " AND COALESCE(p.is_key, 0) = 1 AND COALESCE(p.status, '启用') = '启用'"
    rows = query_rows(
        f"""
        SELECT s.product_code AS 商品编码, s.product_name AS 商品名称,
               SUM(s.quantity) AS 销售数量,
               ROUND(SUM(s.amount), 2) AS 销售金额,
               COUNT(DISTINCT s.salesperson_name) AS 销售业务员数
        FROM sales s
        LEFT JOIN products p ON p.code = s.product_code
        WHERE s.sale_date BETWEEN :start AND :end {extra}
        GROUP BY s.product_code, s.product_name
        ORDER BY 销售数量 DESC
        """,
        sql_params,
    )
    for i, row in enumerate(rows, 1):
        row["排名"] = i
    return rows


def sales_ranking(params):
    selected = params.get("date") or date.today().isoformat()
    selected_month = selected[:7]
    month_start = f"{selected_month}-01"
    sql_params = {"date": selected, "month": selected_month, "month_start": month_start, "month_end": selected}
    extra = ""
    if params.get("salesperson"):
        extra = " WHERE sp.name = :salesperson"
        sql_params["salesperson"] = params["salesperson"]
    rows = query_rows(
        f"""
        SELECT sp.name AS 业务员,
               ROUND(COALESCE(day_sales.amount, 0), 2) AS 今日销售金额,
               ROUND(COALESCE(month_sales.amount, 0), 2) AS 本月销售金额,
               COALESCE(st.peak_target, 0) AS 巅峰目标,
               COALESCE(st.daily_target, 0) AS 每日目标,
               CASE WHEN COALESCE(st.daily_target, 0) = 0 THEN 0
                    ELSE COALESCE(day_sales.amount, 0) / st.daily_target END AS 完成率
        FROM salespeople sp
        LEFT JOIN salesperson_targets st ON st.salesperson_name = sp.name AND st.month = :month
        LEFT JOIN (
            SELECT salesperson_name, SUM(amount) AS amount
            FROM sales
            WHERE sale_date = :date
            GROUP BY salesperson_name
        ) day_sales ON day_sales.salesperson_name = sp.name
        LEFT JOIN (
            SELECT salesperson_name, SUM(amount) AS amount
            FROM sales
            WHERE sale_date BETWEEN :month_start AND :month_end
            GROUP BY salesperson_name
        ) month_sales ON month_sales.salesperson_name = sp.name
        {extra}
        ORDER BY 完成率 DESC, 今日销售金额 DESC
        """,
        sql_params,
    )
    for row in rows:
        row["完成率"] = round(row["完成率"], 4)
    return rows


def gold_rank(params):
    selected = params.get("date") or date.today().isoformat()
    selected_month = params.get("month") or selected[:7]
    selected = selected if selected.startswith(selected_month) else f"{selected_month}-01"
    days = month_days(selected_month)
    sql_params = {"date": selected, "month": selected_month, "month_start": f"{selected_month}-01", "month_end": selected}
    if params.get("salesperson"):
        salesperson_filter = " WHERE sp.name = :salesperson"
        sql_params["salesperson"] = params["salesperson"]
    else:
        salesperson_filter = ""
    rows = query_rows(
        f"""
        SELECT sp.name AS 业务员,
               COALESCE(gt.target_qty, 0) AS 金砖商品目标,
               CASE WHEN :days = 0 THEN 0 ELSE COALESCE(gt.target_qty, 0) / :days END AS 金砖商品每日目标,
               COALESCE(day_sales.qty, 0) AS 每日完成件数,
               CASE WHEN COALESCE(gt.target_qty, 0) = 0 THEN 0
                    ELSE COALESCE(day_sales.qty, 0) / (gt.target_qty / :days) END AS 今日完成率,
               COALESCE(month_sales.qty, 0) AS 累计完成,
               CASE WHEN COALESCE(gt.target_qty, 0) = 0 THEN 0
                    ELSE COALESCE(month_sales.qty, 0) / gt.target_qty END AS 累计完成率
        FROM salespeople sp
        LEFT JOIN gold_targets gt ON gt.salesperson_name = sp.name AND gt.month = :month
        LEFT JOIN (
            SELECT s.salesperson_name, SUM(s.quantity) AS qty
            FROM sales s
            JOIN products p ON p.code = s.product_code AND p.is_key = 1 AND COALESCE(p.status, '启用') = '启用'
            WHERE s.sale_date = :date
            GROUP BY s.salesperson_name
        ) day_sales ON day_sales.salesperson_name = sp.name
        LEFT JOIN (
            SELECT s.salesperson_name, SUM(s.quantity) AS qty
            FROM sales s
            JOIN products p ON p.code = s.product_code AND p.is_key = 1 AND COALESCE(p.status, '启用') = '启用'
            WHERE s.sale_date BETWEEN :month_start AND :month_end
            GROUP BY s.salesperson_name
        ) month_sales ON month_sales.salesperson_name = sp.name
        {salesperson_filter}
        ORDER BY 今日完成率 DESC, 累计完成率 DESC
        """.replace(":days", str(days)),
        sql_params,
    )
    for row in rows:
        row["金砖商品每日目标"] = round(row["金砖商品每日目标"], 2)
        row["今日完成率"] = round(row["今日完成率"], 4)
        row["累计完成率"] = round(row["累计完成率"], 4)
    return rows


def task_stats(params):
    selected_month = params.get("month") or date.today().strftime("%Y-%m")
    selected_date = params.get("date") or date.today().isoformat()
    start = f"{selected_month}-01"
    if selected_month == date.today().strftime("%Y-%m"):
        end = selected_date
    else:
        year, month = map(int, selected_month.split("-"))
        if month == 12:
            end = (date(year + 1, 1, 1) - timedelta(days=1)).isoformat()
        else:
            end = (date(year, month + 1, 1) - timedelta(days=1)).isoformat()

    extra = ""
    sql_params = {"month": selected_month, "date": selected_date, "start": start, "end": end}
    if params.get("salesperson"):
        extra += " AND t.salesperson_name = :salesperson"
        sql_params["salesperson"] = params["salesperson"]
    if params.get("product"):
        extra += " AND (t.product_code = :product OR t.product_name LIKE :product_like)"
        sql_params["product"] = params["product"]
        sql_params["product_like"] = f"%{params['product']}%"
    rows = query_rows(
        f"""
        SELECT t.month AS 月份, t.salesperson_name AS 业务员,
               t.product_code AS 商品编码, t.product_name AS 商品名称,
               COALESCE(day_sales.qty, 0) AS 当日销量,
               COALESCE(month_sales.qty, 0) AS 月累计销量,
               t.target_qty AS 月任务数量,
               CASE WHEN t.target_qty = 0 THEN 0 ELSE COALESCE(month_sales.qty, 0) / t.target_qty END AS 完成率,
               MAX(t.target_qty - COALESCE(month_sales.qty, 0), 0) AS 剩余任务数量,
               CASE WHEN COALESCE(month_sales.qty, 0) >= t.target_qty THEN '已达标' ELSE '未达标' END AS 是否达标
        FROM targets t
        JOIN products p ON p.code = t.product_code AND p.is_key = 1 AND COALESCE(p.status, '启用') = '启用'
        LEFT JOIN (
            SELECT salesperson_name, product_code, SUM(quantity) AS qty
            FROM sales
            WHERE sale_date = :date
            GROUP BY salesperson_name, product_code
        ) day_sales ON day_sales.salesperson_name = t.salesperson_name AND day_sales.product_code = t.product_code
        LEFT JOIN (
            SELECT salesperson_name, product_code, SUM(quantity) AS qty
            FROM sales
            WHERE sale_date BETWEEN :start AND :end
            GROUP BY salesperson_name, product_code
        ) month_sales ON month_sales.salesperson_name = t.salesperson_name AND month_sales.product_code = t.product_code
        WHERE t.month = :month {extra}
        ORDER BY 是否达标 DESC, 完成率 ASC, t.salesperson_name, t.product_code
        """,
        sql_params,
    )
    if params.get("unmetOnly") == "1":
        rows = [r for r in rows if r["是否达标"] == "未达标"]
    for row in rows:
        row["完成率"] = round(row["完成率"], 4)
        row["剩余任务数量"] = max(round(row["剩余任务数量"], 2), 0)
    return rows


def task_rank(params):
    rows = task_stats(params)
    grouped = {}
    for row in rows:
        g = grouped.setdefault(
            row["业务员"],
            {"业务员": row["业务员"], "金砖商品任务数": 0, "已达标商品数": 0, "未达标商品数": 0, "rate_sum": 0, "未达标商品名称": []},
        )
        g["金砖商品任务数"] += 1
        g["rate_sum"] += row["完成率"]
        if row["是否达标"] == "已达标":
            g["已达标商品数"] += 1
        else:
            g["未达标商品数"] += 1
            g["未达标商品名称"].append(row["商品名称"])
    result = []
    for item in grouped.values():
        count = item["金砖商品任务数"] or 1
        item["平均完成率"] = round(item.pop("rate_sum") / count, 4)
        item["未达标商品名称"] = "、".join(item["未达标商品名称"][:8])
        result.append(item)
    return sorted(result, key=lambda x: x["平均完成率"], reverse=True)


def add_pending_item(conn, item_type, ref_key, name, payload):
    exists = conn.execute(
        "SELECT id FROM pending_items WHERE item_type = ? AND ref_key = ? AND status = '待确认'",
        (item_type, ref_key),
    ).fetchone()
    if exists:
        return
    conn.execute(
        """
        INSERT INTO pending_items(item_type, ref_key, name, payload, status, created_at)
        VALUES (?, ?, ?, ?, '待确认', ?)
        """,
        (item_type, ref_key, name, json.dumps(payload, ensure_ascii=False), datetime.now().isoformat(timespec="seconds")),
    )


def ensure_upload_entities(conn, rows):
    new_salespeople = []
    new_products = []
    existing_salespeople = {r["name"] for r in conn.execute("SELECT name FROM salespeople")}
    existing_products = {r["code"] for r in conn.execute("SELECT code FROM products")}

    for row in rows:
        salesperson = row["salesperson_name"]
        if salesperson and salesperson not in existing_salespeople:
            conn.execute(
                "INSERT INTO salespeople(name, region, active, pending_confirm) VALUES (?, '', 1, 1)",
                (salesperson,),
            )
            existing_salespeople.add(salesperson)
            new_salespeople.append(salesperson)
            add_pending_item(conn, "业务员", salesperson, salesperson, {"业务员姓名": salesperson})

        code = row["product_code"]
        name = row["product_name"]
        if code and code not in existing_products:
            conn.execute(
                """
                INSERT INTO products(code, name, category, is_key, status, pending_confirm)
                VALUES (?, ?, '', 0, '启用', 1)
                """,
                (code, name),
            )
            existing_products.add(code)
            new_products.append(code)
            add_pending_item(conn, "商品", code, name, {"商品编码": code, "商品名称": name})

    return new_salespeople, new_products


def dashboard(params):
    selected = params.get("date") or date.today().isoformat()
    overview = one(
        """
        SELECT ROUND(COALESCE(SUM(amount), 0), 2) AS total_amount,
               COALESCE(SUM(quantity), 0) AS total_qty,
               COUNT(DISTINCT customer_name) AS customer_count,
               COUNT(DISTINCT product_code) AS product_count,
               COUNT(DISTINCT salesperson_name) AS salesperson_count
        FROM sales
        WHERE sale_date = :date
        """,
        {"date": selected},
    )
    return {
        "date": selected,
        "overview": overview,
        "salespeopleRank": sales_ranking({"date": selected})[:8],
        "productsRank": product_stats({"period": "day", "date": selected})[:8],
        "taskRank": gold_rank({"month": selected[:7], "date": selected})[:8],
        "unmet": [r for r in gold_rank({"month": selected[:7], "date": selected}) if r["累计完成率"] < 1][:10],
    }


def maintenance_data():
    return {
        "salespeople": query_rows("SELECT id, code, name, region, department, entry_date, active, pending_confirm FROM salespeople ORDER BY id"),
        "products": query_rows("SELECT id, code, name, spec, category, is_key, gold_target, gold_daily_target, status, pending_confirm FROM products ORDER BY code"),
        "targets": query_rows("SELECT id, month, salesperson_name, product_code, product_name, target_qty FROM targets ORDER BY month DESC, salesperson_name, product_code"),
        "salesTargets": query_rows("SELECT id, month, salesperson_name, peak_target, daily_target FROM salesperson_targets ORDER BY month DESC, salesperson_name"),
        "goldTargets": query_rows("SELECT id, code, name, gold_target, gold_daily_target FROM products WHERE is_key = 1 ORDER BY code"),
        "pendingItems": query_rows("SELECT id, item_type, ref_key, name, payload, status, created_at FROM pending_items ORDER BY id DESC"),
    }


class Handler(SimpleHTTPRequestHandler):
    def translate_path(self, path):
        parsed_path = unquote(urlparse(path).path)
        if parsed_path == "/" or parsed_path.startswith("/static/"):
            rel = "static/index.html" if parsed_path == "/" else parsed_path.lstrip("/")
            return str(BASE_DIR / rel)
        return str(BASE_DIR / parsed_path.lstrip("/"))

    def end_headers(self):
        self.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
        self.send_header("Pragma", "no-cache")
        super().end_headers()

    def send_json(self, data, status=200):
        payload = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def do_GET(self):
        parsed = urlparse(self.path)
        params = filters_from(parsed.query)
        try:
            if parsed.path == "/api/dashboard":
                self.send_json(dashboard(params))
            elif parsed.path == "/api/sales-ranking":
                self.send_json(sales_ranking(params))
            elif parsed.path == "/api/gold-rank":
                self.send_json(gold_rank(params))
            elif parsed.path == "/api/salespeople-stats":
                self.send_json(salesperson_stats(params))
            elif parsed.path == "/api/product-stats":
                self.send_json(product_stats(params))
            elif parsed.path == "/api/task-stats":
                self.send_json(task_stats(params))
            elif parsed.path == "/api/task-rank":
                self.send_json(task_rank(params))
            elif parsed.path == "/api/maintenance":
                self.send_json(maintenance_data())
            elif parsed.path.startswith("/api/base-template/"):
                self.download_base_template(parsed.path.split("/")[-1])
            elif parsed.path.startswith("/api/export-maintenance/"):
                self.export_maintenance(parsed.path.split("/")[-1])
            elif parsed.path.startswith("/api/export/"):
                self.export_report(parsed.path.split("/")[-1], params)
            else:
                super().do_GET()
        except Exception as exc:
            traceback.print_exc()
            self.send_json({"error": str(exc)}, 500)

    def do_POST(self):
        parsed = urlparse(self.path)
        try:
            if parsed.path == "/api/upload":
                self.upload_sales(filters_from(parsed.query))
            elif parsed.path == "/api/maintenance":
                length = int(self.headers.get("Content-Length", 0))
                payload = json.loads(self.rfile.read(length).decode("utf-8"))
                self.save_maintenance(payload)
            elif parsed.path == "/api/pending/action":
                length = int(self.headers.get("Content-Length", 0))
                payload = json.loads(self.rfile.read(length).decode("utf-8"))
                self.pending_action(payload)
            elif parsed.path.startswith("/api/import-maintenance/"):
                self.import_maintenance(parsed.path.split("/")[-1])
            elif parsed.path == "/api/reset-demo":
                with db() as conn:
                    seed_demo_data(conn)
                self.send_json({"ok": True, "message": "已重置为模拟数据"})
            else:
                self.send_error(HTTPStatus.NOT_FOUND)
        except Exception as exc:
            traceback.print_exc()
            self.send_json({"error": str(exc)}, 500)

    def upload_sales(self, params):
        form = cgi.FieldStorage(fp=self.rfile, headers=self.headers, environ={"REQUEST_METHOD": "POST"})
        file_item = form["file"] if "file" in form else None
        if file_item is None or not file_item.filename:
            self.send_json({"error": "请选择 .xlsx 文件"}, 400)
            return
        content = file_item.file.read()
        rows, errors = read_sales_excel(content)
        if errors:
            self.send_json({"error": "Excel 校验失败", "errors": errors[:100]}, 400)
            return
        if not rows:
            self.send_json({"error": "Excel 中没有可导入的数据行"}, 400)
            return
        mode = params.get("mode", "")
        upload_dates = sorted(set(row["sale_date"] for row in rows))
        with db() as conn:
            existing = conn.execute(
                f"SELECT sale_date, COUNT(*) AS c FROM sales WHERE sale_date IN ({','.join('?' for _ in upload_dates)}) GROUP BY sale_date",
                upload_dates,
            ).fetchall()
            if existing and mode not in ("merge", "replace"):
                self.send_json(
                    {
                        "duplicate": True,
                        "message": "检测到相同日期已有销售数据，请选择覆盖或合并。",
                        "dates": [dict(r) for r in existing],
                    },
                    409,
                )
                return
            if mode == "replace" and upload_dates:
                conn.execute(
                    f"DELETE FROM sales WHERE sale_date IN ({','.join('?' for _ in upload_dates)})",
                    upload_dates,
                )
            new_salespeople, new_products = ensure_upload_entities(conn, rows)
            now = datetime.now().isoformat(timespec="seconds")
            for row in rows:
                conn.execute(
                    """
                    INSERT INTO sales(
                        sale_date, salesperson_name, customer_name, product_code, product_name,
                        quantity, unit_price, amount, uploaded_at, source_file
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        row["sale_date"],
                        row["salesperson_name"],
                        row["customer_name"],
                        row["product_code"],
                        row["product_name"],
                        row["quantity"],
                        row["unit_price"],
                        row["amount"],
                        now,
                        file_item.filename,
                    ),
                )
            conn.commit()
        message = f"成功导入 {len(rows)} 条销售明细"
        if new_salespeople or new_products:
            message += f"，自动识别新业务员 {len(new_salespeople)} 个、新商品 {len(new_products)} 个，已加入待确认数据池"
        self.send_json(
            {
                "ok": True,
                "rows": len(rows),
                "dates": upload_dates,
                "newSalespeople": new_salespeople,
                "newProducts": new_products,
                "message": message,
            }
        )

    def send_xlsx_file(self, output, filename):
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        safe_name = f"download_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        self.send_header("Content-Disposition", f"attachment; filename=\"{safe_name}\"; filename*=UTF-8''{quote(filename)}")
        self.send_header("Content-Length", str(output.stat().st_size))
        self.end_headers()
        self.wfile.write(output.read_bytes())

    def download_base_template(self, table):
        config = BASE_IMPORT_CONFIG.get(table)
        if not config:
            self.send_json({"error": "未知模板类型"}, 404)
            return
        export_base_templates()
        filename = f"{config['title']}导入模板.xlsx"
        self.send_xlsx_file(EXPORT_DIR / filename, filename)

    def export_maintenance(self, table):
        data = maintenance_data()
        if table == "salespeople":
            headers = BASE_IMPORT_CONFIG[table]["headers"]
            rows = [
                {
                    "业务员编码": r.get("code") or "",
                    "业务员姓名": r.get("name") or "",
                    "所属区域": r.get("region") or "",
                    "所属部门": r.get("department") or "",
                    "入职日期": r.get("entry_date") or "",
                    "状态": "启用" if r.get("active") else "停用",
                }
                for r in data["salespeople"]
            ]
        elif table == "products":
            headers = BASE_IMPORT_CONFIG[table]["headers"]
            rows = [
                {
                    "商品编码": r.get("code") or "",
                    "商品名称": r.get("name") or "",
                    "商品规格": r.get("spec") or "",
                    "商品分类": r.get("category") or "",
                    "是否重点商品": "是" if r.get("is_key") else "否",
                    "状态": r.get("status") or "启用",
                }
                for r in data["products"]
            ]
        elif table == "goldProducts":
            headers = BASE_IMPORT_CONFIG[table]["headers"]
            rows = [
                {
                    "商品编码": r.get("code") or "",
                    "商品名称": r.get("name") or "",
                    "金砖商品目标": r.get("gold_target") or 0,
                    "每日目标": r.get("gold_daily_target") or 0,
                }
                for r in data["goldTargets"]
            ]
        elif table == "salesTargets":
            headers = BASE_IMPORT_CONFIG[table]["headers"]
            rows = [
                {
                    "业务员": r.get("salesperson_name") or "",
                    "巅峰目标": r.get("peak_target") or 0,
                    "每日目标": r.get("daily_target") or 0,
                }
                for r in data["salesTargets"]
            ]
        else:
            self.send_json({"error": "未知导出类型"}, 404)
            return
        filename = f"{BASE_IMPORT_CONFIG[table]['title']}_当前数据_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output = EXPORT_DIR / filename
        write_xlsx(output, headers, rows)
        self.send_xlsx_file(output, filename)

    def import_maintenance(self, table):
        config = BASE_IMPORT_CONFIG.get(table)
        if not config:
            self.send_json({"error": "未知导入类型"}, 404)
            return
        form = cgi.FieldStorage(fp=self.rfile, headers=self.headers, environ={"REQUEST_METHOD": "POST"})
        file_item = form["file"] if "file" in form else None
        if file_item is None or not file_item.filename:
            self.send_json({"error": "请选择 Excel 文件"}, 400)
            return
        rows, read_errors = read_excel_dicts(file_item.file.read(), file_item.filename, config["headers"])
        errors = [{"行号": "", "错误原因": err} for err in read_errors]
        success = 0
        if not read_errors:
            with db() as conn:
                for row in rows:
                    row_no = row.pop("_row_no")
                    try:
                        if table == "salespeople":
                            name = str(row.get("业务员姓名") or "").strip()
                            if not name:
                                raise ValueError("业务员姓名不能为空")
                            conn.execute(
                                """
                                INSERT INTO salespeople(code, name, region, department, entry_date, active, pending_confirm)
                                VALUES (?, ?, ?, ?, ?, ?, 0)
                                ON CONFLICT(name) DO UPDATE SET
                                    code=excluded.code, region=excluded.region, department=excluded.department,
                                    entry_date=excluded.entry_date, active=excluded.active
                                """,
                                (
                                    str(row.get("业务员编码") or "").strip(),
                                    name,
                                    str(row.get("所属区域") or "").strip(),
                                    str(row.get("所属部门") or "").strip(),
                                    parse_date(row.get("入职日期")) or str(row.get("入职日期") or "").strip(),
                                    status_active(row.get("状态")),
                                ),
                            )
                        elif table == "products":
                            code = str(row.get("商品编码") or "").strip()
                            name = str(row.get("商品名称") or "").strip()
                            if not code or not name:
                                raise ValueError("商品编码、商品名称不能为空")
                            conn.execute(
                                """
                                INSERT INTO products(code, name, spec, category, is_key, status, pending_confirm)
                                VALUES (?, ?, ?, ?, ?, ?, 0)
                                ON CONFLICT(code) DO UPDATE SET
                                    name=excluded.name, spec=excluded.spec, category=excluded.category,
                                    is_key=excluded.is_key, status=excluded.status
                                """,
                                (
                                    code,
                                    name,
                                    str(row.get("商品规格") or "").strip(),
                                    str(row.get("商品分类") or "").strip(),
                                    1 if truthy(row.get("是否重点商品")) else 0,
                                    str(row.get("状态") or "启用").strip() or "启用",
                                ),
                            )
                        elif table == "goldProducts":
                            code = str(row.get("商品编码") or "").strip()
                            name = str(row.get("商品名称") or "").strip()
                            if not code or not name:
                                raise ValueError("商品编码、商品名称不能为空")
                            conn.execute(
                                """
                                INSERT INTO products(code, name, spec, category, is_key, gold_target, gold_daily_target, status, pending_confirm)
                                VALUES (?, ?, '', '', 1, ?, ?, '启用', 0)
                                ON CONFLICT(code) DO UPDATE SET
                                    name=excluded.name, is_key=1,
                                    gold_target=excluded.gold_target,
                                    gold_daily_target=excluded.gold_daily_target
                                """,
                                (code, name, to_float(row.get("金砖商品目标")), to_float(row.get("每日目标"))),
                            )
                        elif table == "salesTargets":
                            salesperson = str(row.get("业务员") or "").strip()
                            if not salesperson:
                                raise ValueError("业务员不能为空")
                            current_month = date.today().strftime("%Y-%m")
                            conn.execute(
                                """
                                INSERT INTO salesperson_targets(month, salesperson_name, peak_target, daily_target)
                                VALUES (?, ?, ?, ?)
                                ON CONFLICT(month, salesperson_name) DO UPDATE SET
                                    peak_target=excluded.peak_target,
                                    daily_target=excluded.daily_target
                                """,
                                (current_month, salesperson, to_float(row.get("巅峰目标")), to_float(row.get("每日目标"))),
                            )
                        success += 1
                    except Exception as exc:
                        errors.append({"行号": row_no, "错误原因": str(exc)})
                conn.commit()
        error_url = ""
        if errors:
            filename = f"{config['title']}导入错误明细_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            write_xlsx(EXPORT_DIR / filename, ["行号", "错误原因"], errors)
            error_url = f"/exports/{filename}"
        self.send_json({"ok": True, "success": success, "failed": len(errors), "errorUrl": error_url})

    def save_maintenance(self, payload):
        table = payload.get("table")
        records = payload.get("records", [])
        with db() as conn:
            if table == "salespeople":
                conn.execute("DELETE FROM salespeople")
                for r in records:
                    conn.execute(
                        """
                        INSERT INTO salespeople(code, name, region, department, entry_date, active, pending_confirm)
                        VALUES (?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            r.get("code", "").strip(),
                            r.get("name", "").strip(),
                            r.get("region", "").strip(),
                            r.get("department", "").strip(),
                            r.get("entry_date", "").strip(),
                            1 if r.get("active") else 0,
                            1 if r.get("pending_confirm") else 0,
                        ),
                    )
            elif table == "products":
                conn.execute("DELETE FROM products")
                for r in records:
                    conn.execute(
                        """
                        INSERT INTO products(code, name, spec, category, is_key, gold_target, gold_daily_target, status, pending_confirm)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            r.get("code", "").strip(),
                            r.get("name", "").strip(),
                            r.get("spec", "").strip(),
                            r.get("category", "").strip(),
                            1 if r.get("is_key") else 0,
                            to_float(r.get("gold_target")),
                            to_float(r.get("gold_daily_target")),
                            r.get("status", "启用").strip() or "启用",
                            1 if r.get("pending_confirm") else 0,
                        ),
                    )
            elif table == "targets":
                conn.execute("DELETE FROM targets")
                for r in records:
                    conn.execute(
                        """
                        INSERT INTO targets(month, salesperson_name, product_code, product_name, target_qty)
                        VALUES (?, ?, ?, ?, ?)
                        """,
                        (
                            r.get("month", "").strip(),
                            r.get("salesperson_name", "").strip(),
                            r.get("product_code", "").strip(),
                            r.get("product_name", "").strip(),
                            to_float(r.get("target_qty")),
                        ),
                    )
            elif table == "salesTargets":
                conn.execute("DELETE FROM salesperson_targets")
                for r in records:
                    conn.execute(
                        """
                        INSERT INTO salesperson_targets(month, salesperson_name, peak_target, daily_target)
                        VALUES (?, ?, ?, ?)
                        """,
                        (
                            r.get("month", "").strip(),
                            r.get("salesperson_name", "").strip(),
                            to_float(r.get("peak_target")),
                            to_float(r.get("daily_target")),
                        ),
                    )
            elif table == "goldProducts":
                for r in records:
                    conn.execute(
                        """
                        INSERT INTO products(code, name, spec, category, is_key, gold_target, gold_daily_target, status, pending_confirm)
                        VALUES (?, ?, '', '', 1, ?, ?, '启用', 0)
                        ON CONFLICT(code) DO UPDATE SET
                            name=excluded.name,
                            is_key=1,
                            gold_target=excluded.gold_target,
                            gold_daily_target=excluded.gold_daily_target
                        """,
                        (
                            r.get("code", "").strip(),
                            r.get("name", "").strip(),
                            to_float(r.get("gold_target")),
                            to_float(r.get("gold_daily_target")),
                        ),
                    )
            else:
                self.send_json({"error": "未知维护表"}, 400)
                return
            conn.commit()
        self.send_json({"ok": True, "message": "保存成功"})

    def pending_action(self, payload):
        action = payload.get("action")
        item_id = payload.get("id")
        with db() as conn:
            item = conn.execute("SELECT * FROM pending_items WHERE id = ?", (item_id,)).fetchone()
            if not item:
                self.send_json({"error": "待确认数据不存在"}, 404)
                return
            if action == "confirm":
                if item["item_type"] == "业务员":
                    conn.execute("UPDATE salespeople SET pending_confirm = 0 WHERE name = ?", (item["ref_key"],))
                elif item["item_type"] == "商品":
                    conn.execute("UPDATE products SET pending_confirm = 0 WHERE code = ?", (item["ref_key"],))
                conn.execute("UPDATE pending_items SET status = '已确认' WHERE id = ?", (item_id,))
            elif action == "modify":
                new_key = str(payload.get("ref_key") or "").strip()
                new_name = str(payload.get("name") or "").strip()
                if not new_key or not new_name:
                    self.send_json({"error": "编码/姓名和名称不能为空"}, 400)
                    return
                old_key = item["ref_key"]
                if item["item_type"] == "业务员":
                    conn.execute("UPDATE salespeople SET name = ? WHERE name = ?", (new_name, old_key))
                    conn.execute("UPDATE sales SET salesperson_name = ? WHERE salesperson_name = ?", (new_name, old_key))
                    conn.execute("UPDATE salesperson_targets SET salesperson_name = ? WHERE salesperson_name = ?", (new_name, old_key))
                    conn.execute("UPDATE gold_targets SET salesperson_name = ? WHERE salesperson_name = ?", (new_name, old_key))
                    new_key = new_name
                elif item["item_type"] == "商品":
                    conn.execute("UPDATE products SET code = ?, name = ? WHERE code = ?", (new_key, new_name, old_key))
                    conn.execute("UPDATE sales SET product_code = ?, product_name = ? WHERE product_code = ?", (new_key, new_name, old_key))
                    conn.execute("UPDATE targets SET product_code = ?, product_name = ? WHERE product_code = ?", (new_key, new_name, old_key))
                conn.execute(
                    """
                    UPDATE pending_items
                    SET ref_key = ?, name = ?, payload = ?
                    WHERE id = ?
                    """,
                    (
                        new_key,
                        new_name,
                        json.dumps({"修改后编码/姓名": new_key, "修改后名称": new_name}, ensure_ascii=False),
                        item_id,
                    ),
                )
            elif action == "delete":
                if item["item_type"] == "业务员":
                    conn.execute("DELETE FROM sales WHERE salesperson_name = ?", (item["ref_key"],))
                    conn.execute("DELETE FROM salesperson_targets WHERE salesperson_name = ?", (item["ref_key"],))
                    conn.execute("DELETE FROM gold_targets WHERE salesperson_name = ?", (item["ref_key"],))
                    conn.execute("DELETE FROM salespeople WHERE name = ?", (item["ref_key"],))
                elif item["item_type"] == "商品":
                    conn.execute("DELETE FROM sales WHERE product_code = ?", (item["ref_key"],))
                    conn.execute("DELETE FROM targets WHERE product_code = ?", (item["ref_key"],))
                    conn.execute("DELETE FROM products WHERE code = ?", (item["ref_key"],))
                conn.execute("UPDATE pending_items SET status = '已删除' WHERE id = ?", (item_id,))
            else:
                self.send_json({"error": "未知操作"}, 400)
                return
            conn.commit()
        self.send_json({"ok": True, "message": "操作成功"})

    def export_report(self, name, params):
        if name == "sales-ranking":
            rows = sales_ranking(params)
            headers = ["业务员", "今日销售金额", "本月销售金额", "巅峰目标", "每日目标", "完成率"]
            title = "业务销售排行"
            for r in rows:
                r["完成率"] = f"{r['完成率']:.2%}"
        elif name == "gold-rank":
            rows = gold_rank(params)
            headers = ["业务员", "金砖商品目标", "金砖商品每日目标", "每日完成件数", "今日完成率", "累计完成", "累计完成率"]
            title = "金砖商品完成率排行"
            for r in rows:
                r["今日完成率"] = f"{r['今日完成率']:.2%}"
                r["累计完成率"] = f"{r['累计完成率']:.2%}"
        elif name == "salespeople":
            rows = salesperson_stats(params)
            headers = ["排名", "业务员", "销售金额", "销售数量", "成交客户数", "销售商品数"]
            title = "业务员销售统计"
        elif name == "products":
            rows = product_stats(params)
            headers = ["排名", "商品编码", "商品名称", "销售数量", "销售金额", "销售业务员数"]
            title = "商品销售统计"
        elif name == "tasks":
            rows = task_stats(params)
            headers = ["月份", "业务员", "商品编码", "商品名称", "当日销量", "月累计销量", "月任务数量", "完成率", "剩余任务数量", "是否达标"]
            title = "金砖商品任务统计"
            for r in rows:
                r["完成率"] = f"{r['完成率']:.2%}"
        elif name == "task-rank":
            rows = task_rank(params)
            headers = ["业务员", "金砖商品任务数", "已达标商品数", "未达标商品数", "平均完成率", "未达标商品名称"]
            title = "金砖商品完成率排行榜"
            for r in rows:
                r["平均完成率"] = f"{r['平均完成率']:.2%}"
        else:
            self.send_json({"error": "未知导出类型"}, 404)
            return
        filename = f"{title}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output = EXPORT_DIR / filename
        write_xlsx(output, headers, rows)
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        safe_name = f"sales_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        self.send_header("Content-Disposition", f"attachment; filename=\"{safe_name}\"; filename*=UTF-8''{quote(filename)}")
        self.send_header("Content-Length", str(output.stat().st_size))
        self.end_headers()
        self.wfile.write(output.read_bytes())


def main():
    init_db(seed="--reset" in sys.argv)
    port = int(os.environ.get("PORT", "8000"))
    httpd = ThreadingHTTPServer(("127.0.0.1", port), Handler)
    print(f"销售报表自动统计系统 V1 已启动：http://127.0.0.1:{port}")
    print(f"测试上传文件：{EXPORT_DIR / '模拟销售明细_最近30天.xlsx'}")
    httpd.serve_forever()


if __name__ == "__main__":
    main()
